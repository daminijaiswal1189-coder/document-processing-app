"""
Excel processing: append POC Status and highlight FALSE/NA cells.

Flow:
  1. Load workbook from storage/uploads (.xlsx directly, or convert .xls first).
  2. Add a new column at the end with header POC Status.
  3. For each data row (non-summary, has content), set Processed and red-fill FALSE/NA cells.
  4. Add worksheet ``Name and SSN`` with columns name and ssn copied from the source sheet.
  5. Save to storage/processed/processed_<uuid>.xlsx.

Uses worksheet._cells for reads where possible to avoid creating empty cells.
"""

import logging
import os
import tempfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from app.services.excel.excel_config import (
    EXTRACT_SHEET_NAME,
    EXTRACT_TAB_NAME_HEADER,
    EXTRACT_TAB_SSN_HEADER,
    FALSE_NA_FILL_RGB,
    HEADER_ROW,
    NEW_COLUMN_DEFAULT_VALUE,
    NEW_COLUMN_HEADER,
    SOURCE_NAME_HEADERS,
    SOURCE_SSN_HEADERS,
)
from app.services.excel.xls_converter import convert_xls_to_xlsx

logger = logging.getLogger(__name__)

# Reused PatternFill for FALSE / NA highlighting (copied per cell).
_RED_FILL = PatternFill(fill_type="solid", fgColor=FALSE_NA_FILL_RGB)


class ExcelProcessor:
    """
    Transforms an uploaded .xlsx or legacy .xls into a processed .xlsx copy with an extra status column.

    Summary rows (containing 'total' in columns A–E) are skipped for POC values.
    """

    def process(self, source_path: Path, output_path: Path) -> dict[str, str | int]:
        """
        Run the full Excel pipeline and write ``output_path``.

        Args:
            source_path: Path to the uploaded .xlsx or .xls in storage/uploads.
            output_path: Destination .xlsx path under storage/processed.

        Returns:
            Metadata for the API response (sheet name, column index, counts).
        """
        if not source_path.is_file():
            raise FileNotFoundError(f"Source file not found: {source_path}")

        logger.info("Processing Excel file: %s", source_path.name)

        temp_dir: tempfile.TemporaryDirectory[str] | None = None
        work_path = source_path
        if source_path.suffix.lower() == ".xls":
            temp_dir = tempfile.TemporaryDirectory(prefix="poc_xls_")
            work_path = Path(temp_dir.name) / f"{source_path.stem}.xlsx"
            convert_xls_to_xlsx(source_path, work_path)

        try:
            workbook = load_workbook(work_path, data_only=False)
            worksheet = workbook.active
            result = self._apply_transforms(workbook, worksheet)

            output_path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(output_path)
            workbook.close()

            logger.info(
                "Saved processed file: %s (rows updated: %s, new column: %s, red fills: %s)",
                output_path.name,
                result["rows_updated"],
                result["new_column_index"],
                result["cells_filled_red"],
            )

            return result
        finally:
            if temp_dir is not None:
                temp_dir.cleanup()

    def process_inplace(self, target_path: Path) -> dict[str, str | int]:
        """
        Apply the same transforms as ``process`` and overwrite ``target_path`` in place.

        Uses a temporary file in the same directory, then atomic replace, to reduce
        risk of a partial write if save fails.

        Args:
            target_path: Existing .xlsx on disk (must be writable).

        Returns:
            Transform metadata (sheet name, column index, counts).
        """
        if not target_path.is_file():
            raise FileNotFoundError(f"Target file not found: {target_path}")
        if target_path.suffix.lower() != ".xlsx":
            raise ValueError("In-place processing supports .xlsx only")

        logger.info("In-place Excel update: %s", target_path)

        workbook = load_workbook(target_path, data_only=False)
        try:
            worksheet = workbook.active
            result = self._apply_transforms(workbook, worksheet)

            fd, tmp_name = tempfile.mkstemp(
                suffix=".xlsx",
                prefix=f".{target_path.stem}_poc_",
                dir=target_path.parent,
            )
            os.close(fd)
            tmp_path = Path(tmp_name)
            try:
                workbook.save(tmp_path)
                os.replace(tmp_path, target_path)
            finally:
                if tmp_path.is_file():
                    tmp_path.unlink(missing_ok=True)

            logger.info(
                "Updated in place: %s (rows updated: %s, column: %s)",
                target_path.name,
                result["rows_updated"],
                result["new_column_index"],
            )
            return result
        finally:
            workbook.close()

    def _apply_transforms(self, workbook, worksheet) -> dict[str, str | int]:
        """POC Status column, FALSE/NA highlighting, and Name/SSN extract sheet."""
        new_col = self._resolve_or_create_poc_column(worksheet)

        rows_updated = 0
        cells_filled_red = 0
        last_data_row = self._find_last_data_row(worksheet)
        for row in range(HEADER_ROW + 1, last_data_row + 1):
            if not self._row_should_receive_value(worksheet, row):
                continue
            worksheet.cell(row=row, column=new_col, value=NEW_COLUMN_DEFAULT_VALUE)
            cells_filled_red += self._fill_color_in_cell(worksheet, row)
            rows_updated += 1

        name_ssn_rows = self._create_name_ssn_sheet(workbook, worksheet)
        sheet_name = worksheet.title

        return {
            "sheet_name": sheet_name,
            "new_column_index": new_col,
            "new_column_header": NEW_COLUMN_HEADER,
            "rows_updated": rows_updated,
            "cells_filled_red": cells_filled_red,
            "extract_sheet_name": EXTRACT_SHEET_NAME,
            "name_ssn_rows": name_ssn_rows,
        }

    def _resolve_or_create_poc_column(self, worksheet) -> int:
        """Use existing POC Status column if present; otherwise append a new column."""
        for col in range(1, worksheet.max_column + 1):
            header = self._get_cell_value(worksheet, HEADER_ROW, col)
            if header is None:
                continue
            if str(header).strip() == NEW_COLUMN_HEADER:
                return col
        new_col = worksheet.max_column + 1
        worksheet.cell(row=HEADER_ROW, column=new_col, value=NEW_COLUMN_HEADER)
        return new_col

    def _fill_color_in_cell(self, worksheet, row: int) -> int:
        """
        Apply red fill to existing cells on ``row`` whose value is FALSE or NA.

        Only touches cells already present in worksheet._cells.

        Returns:
            Number of cells that received the red fill.
        """
        filled = 0
        for (cell_row, _col), cell in worksheet._cells.items():
            if cell_row != row:
                continue
            if not self._is_false_or_na_value(cell.value):
                continue
            cell.fill = copy(_RED_FILL)
            filled += 1
        return filled

    def _is_false_or_na_value(self, value) -> bool:
        """True for boolean False or string false / na / n/a (case-insensitive)."""
        if isinstance(value, bool):
            return value is False
        if value is None:
            return False
        text = str(value).strip().lower()
        return text in {"false", "na", "n/a"}

    def _find_last_data_row(self, worksheet) -> int:
        """Last row index that contains any formula or non-empty value (scan existing cells)."""
        last_row = HEADER_ROW
        for (row, _col), cell in worksheet._cells.items():
            if row <= HEADER_ROW:
                continue
            if self._cell_has_meaningful_content(cell):
                last_row = max(last_row, row)
        return max(last_row, HEADER_ROW + 1)

    def _cell_has_meaningful_content(self, cell) -> bool:
        """True if the cell has a formula or non-blank scalar value."""
        if cell.data_type == "f" and cell.value:
            return True
        if cell.value is None:
            return False
        return str(cell.value).strip() != ""

    def _get_cell_value(self, worksheet, row: int, col: int):
        """Read cell value without creating a new cell if missing."""
        cell = worksheet._cells.get((row, col))
        if cell is None:
            return None
        return cell.value

    def _row_should_receive_value(self, worksheet, row: int) -> bool:
        """
        Decide whether this row gets POC Status = Processed.

        Skips summary rows; requires at least one meaningful cell on the row.
        """
        if self._is_summary_row(worksheet, row):
            return False
        return self._row_has_meaningful_content(worksheet, row)

    def _row_has_meaningful_content(self, worksheet, row: int) -> bool:
        """True if any existing cell on ``row`` has meaningful content."""
        for (_row, _col), cell in worksheet._cells.items():
            if _row != row:
                continue
            if self._cell_has_meaningful_content(cell):
                return True
        return False

    def _is_summary_row(self, worksheet, row: int) -> bool:
        """True when columns A–E contain a label with 'total' (case-insensitive)."""
        for col in range(1, min(worksheet.max_column, 5) + 1):
            value = self._get_cell_value(worksheet, row, col)
            if value is None:
                continue
            if "total" in str(value).strip().lower():
                return True
        return False

    def _find_header_column(self, worksheet, header_names: tuple[str, ...]) -> int | None:
        """Return 1-based column index whose row-1 header matches one of ``header_names``."""
        needles = {name.strip().lower() for name in header_names}
        for col in range(1, worksheet.max_column + 1):
            value = self._get_cell_value(worksheet, HEADER_ROW, col)
            if value is None:
                continue
            if str(value).strip().lower() in needles:
                return col
        return None

    def _create_name_ssn_sheet(self, workbook, worksheet) -> int:
        """
        Add a tab with columns ``name`` and ``ssn`` populated from the active sheet.

        Skips summary rows and rows without a Name value.

        Returns:
            Number of data rows written (excluding header).
        """
        name_col = self._find_header_column(worksheet, SOURCE_NAME_HEADERS)
        if name_col is None:
            raise ValueError(
                f'Name column not found in row {HEADER_ROW} (expected header: Name)'
            )
        ssn_col = self._find_header_column(worksheet, SOURCE_SSN_HEADERS)

        if EXTRACT_SHEET_NAME in workbook.sheetnames:
            del workbook[EXTRACT_SHEET_NAME]

        extract_ws = workbook.create_sheet(title=EXTRACT_SHEET_NAME[:31])
        extract_ws.cell(row=HEADER_ROW, column=1, value=EXTRACT_TAB_NAME_HEADER)
        extract_ws.cell(row=HEADER_ROW, column=2, value=EXTRACT_TAB_SSN_HEADER)

        out_row = HEADER_ROW + 1
        last_data_row = self._find_last_data_row(worksheet)
        for row in range(HEADER_ROW + 1, last_data_row + 1):
            if self._is_summary_row(worksheet, row):
                continue
            name_val = self._get_cell_value(worksheet, row, name_col)
            if name_val is None or str(name_val).strip() == "":
                continue
            ssn_val = self._get_cell_value(worksheet, row, ssn_col) if ssn_col else None
            extract_ws.cell(row=out_row, column=1, value=name_val)
            extract_ws.cell(row=out_row, column=2, value=ssn_val)
            out_row += 1

        rows_written = out_row - HEADER_ROW - 1
        logger.info(
            "Created sheet %r with %s name/ssn rows (name col=%s, ssn col=%s)",
            EXTRACT_SHEET_NAME,
            rows_written,
            name_col,
            ssn_col,
        )
        return rows_written
