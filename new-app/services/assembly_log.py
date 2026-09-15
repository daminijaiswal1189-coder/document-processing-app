from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from models.plan_profile import PlanProfile

# Date/Time are when Process ran (assembly), not when the tester was emailed.
_HEADERS = ["Date", "Time", "Plan Number", "Plan Name", "PYE", "Filename", "Job ID"]


def append_row(log_path: Path, profile: PlanProfile, filename: str, job_id: str) -> Path:
    """TEST23 §P: append the assembled package to the Val Assembly Log.

    Date and Time are the local assembly timestamp. The app does not record
    a separate 'date sent' because it only writes an .eml draft.
    """
    log_path.parent.mkdir(parents=True, exist_ok=True)
    if log_path.is_file():
        book = load_workbook(log_path)
        sheet = book.active
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "Val Assembly Log"
        for col, header in enumerate(_HEADERS, start=1):
            sheet.cell(1, col, header)
    now = datetime.now()
    sheet.append(
        [
            now.strftime("%m/%d/%Y"),
            now.strftime("%H:%M"),
            profile.plan_number or "",
            profile.plan_name or "",
            profile.plan_year_end or "",
            filename,
            job_id,
        ]
    )
    book.save(log_path)
    return log_path
