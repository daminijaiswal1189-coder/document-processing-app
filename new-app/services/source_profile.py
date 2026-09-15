from __future__ import annotations

import io
import re
from datetime import datetime
from typing import Any

from openpyxl import load_workbook

from models.plan_profile import PlanProfile

_PLAN_NUMBER = re.compile(r"\b(\d{5,7})\b")
_DATE = re.compile(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})")
_MONTH_DATE = re.compile(
    r"(January|February|March|April|May|June|July|August|September|October|November|December)"
    r"\s+(\d{1,2}),\s+(\d{4})",
    re.I,
)
_PERCENT = re.compile(r"(\d+(?:\.\d+)?)\s*%")


def parse_excel(data: bytes) -> dict[str, Any]:
    """Read the Val Assembly / Omnium checklist (CURRENT vs PRIOR, reclass, HCE)."""
    book = load_workbook(io.BytesIO(data), data_only=True)
    sheet = book.active
    cells: dict[str, str] = {}
    headers: list[str] = []
    header_row = 1
    for col in range(1, (sheet.max_column or 1) + 1):
        headers.append(_cell_str(sheet.cell(header_row, col).value))
    if headers and any("plan number" in h.lower() for h in headers):
        for col, header in enumerate(headers, start=1):
            value = sheet.cell(2, col).value
            if header:
                cells[header.lower()] = _cell_str(value)
    for row in sheet.iter_rows(min_row=5, max_row=sheet.max_row or 5, max_col=8, values_only=True):
        label = _cell_str(row[0] if row else None)
        value = _cell_str(row[1] if row and len(row) > 1 else None)
        if label:
            cells[label.lower()] = value

    method = None
    for key, value in cells.items():
        if "prior" in key and "current" in key:
            method = _method(value)
            break
    if method is None:
        method = _method(_first(cells, "testing method", "adp/acp method"))
    reclass = None
    for key, value in cells.items():
        if "reclassif" in key or "shifting" in key:
            reclass = _yes_no(value)
            break
    top_heavy = _yes_no(_first(cells, "top heavy"))
    if top_heavy is None and "403b" in _first(cells, "top heavy").lower():
        top_heavy = False

    pye = _normalize_date(
        _first(cells, "pye", "plan year end") or _first(cells, "plan year")
    )
    overlay: dict[str, Any] = {
        "plan_number": _digits(_first(cells, "omni plan number", "plan number", "plan #")),
        "plan_name": _first(cells, "plan name") or None,
        "plan_type": _plan_type(_first(cells, "plan type")),
        "plan_year_end": pye,
        "testing_method": method,
        "adp_reclassified": reclass,
        "hce_current": _yes_no(_first(cells, "hce current", "hce currently")),
        "hce_future": _yes_no(_first(cells, "hce future")),
        "per_payroll_match": _yes_no(_first(cells, "per payroll match", "per payroll")),
    }
    if method is None and "sh" in _first(cells, "prior", "current", "adp/acp").lower():
        overlay["safe_harbor"] = True
    notes = " ".join(value for key, value in cells.items() if "note" in key or "comment" in key)
    if re.search(r"variance", notes, re.I):
        overlay["variance_report"] = True
    if re.search(r"contribution and/or adjustment|census adjustment|contributions required", notes, re.I):
        overlay["contributions_required"] = True
    if re.search(r"safe harbor|\bSH\b", notes, re.I):
        overlay["safe_harbor"] = True
    if re.search(r"catch-?up.{0,40}not allowed", notes, re.I):
        overlay["catchup_disallowed"] = True
    if re.search(r"partially vested", notes, re.I):
        overlay["partially_vested"] = True
    if top_heavy is not None:
        overlay["top_heavy"] = top_heavy
    if reclass is True:
        overlay["returns_required"] = False
    return {key: value for key, value in overlay.items() if value not in (None, "")}


def parse_summary_text(text: str) -> dict[str, Any]:
    """Read Compliance Testing Summary of Results pass/fail lines."""
    blob = re.sub(r"\s+", " ", text or "")
    overlay: dict[str, Any] = {}
    numbers = _PLAN_NUMBER.findall(text or "")
    if numbers:
        overlay["plan_number"] = numbers[0]
    name = _summary_plan_name(text or "")
    if name:
        overlay["plan_name"] = name
    pye = _summary_pye(text or "")
    if pye:
        overlay["plan_year_end"] = pye

    percent = _PERCENT.search(text or "")
    if percent and re.search(r"top\s*heavy", text or "", re.I):
        overlay["top_heavy_percent"] = float(percent.group(1))
    if re.search(r"plan is not top heavy", text or "", re.I):
        overlay["top_heavy"] = False
    elif re.search(r"plan is top heavy", text or "", re.I):
        overlay["top_heavy"] = True

    overlay["fail_402g"] = _tested_fail(blob, r"402\s*\(?g\)?")
    overlay["fail_415"] = _tested_fail(blob, r"415\s*\(?c\)?|Annual Additions")
    overlay["adp_failed"] = _tested_fail(blob, r"Average Deferral Percentage|\bADP\b")
    overlay["acp_failed"] = _tested_fail(blob, r"Average Contribution Percentage|\bACP\b")
    if overlay["acp_failed"] is None and not re.search(r"\bACP\b", blob, re.I):
        overlay.pop("acp_failed")

    if re.search(r"already been processed", blob, re.I):
        overlay["returns_already_processed"] = True
        overlay["returns_required"] = False
    elif overlay.get("adp_failed") or overlay.get("acp_failed") or overlay.get("fail_402g") or overlay.get("fail_415"):
        if re.search(r"excess returns|returns are required|refunds", blob, re.I):
            overlay["returns_required"] = True
    if "compliance testing summary of results" in blob.lower():
        overlay["after_12_months"] = bool(re.search(r"after\s+12[- ]month", blob, re.I))
        overlay["variance_report"] = bool(re.search(r"(?<!no )variance report", blob, re.I))
    return {key: value for key, value in overlay.items() if value not in (None, "")}


def merge_overlays(profile: PlanProfile, overlays: list[dict[str, Any]], warnings: list[str]) -> PlanProfile:
    numbers = [str(item["plan_number"]) for item in overlays if item.get("plan_number")]
    if profile.plan_number:
        numbers.append(profile.plan_number)
    unique = list(dict.fromkeys(numbers))
    if len(unique) > 1:
        warnings.append(
            "Plan number differs between Excel and summary (" + ", ".join(unique) + "). "
            "Pass/fail comes from the Compliance Testing Summary."
        )
    for overlay in overlays:
        for key, value in overlay.items():
            setattr(profile, key, value)
    return profile


def _tested_fail(blob: str, name: str) -> bool | None:
    match = re.search(r"(?:" + name + r").{0,50}?Tested:\s*(Pass|Fail)", blob, flags=re.I)
    if not match:
        return None
    return match.group(1).lower() == "fail"


def _summary_plan_name(text: str) -> str | None:
    for line in text.splitlines():
        cleaned = re.sub(r"\s+", " ", line).strip()
        if re.search(r"401\s*\(?k\)?|403\s*\(?b\)?|savings plan", cleaned, re.I) and not re.search(
            r"summary of results|plan year", cleaned, re.I
        ):
            return cleaned
    return None


def _summary_pye(text: str) -> str | None:
    match = re.search(r"Plan Year End:\s*(.+)", text, flags=re.I)
    raw = match.group(1).strip() if match else ""
    month = _MONTH_DATE.search(raw or text)
    if month:
        parsed = datetime.strptime(
            f"{month.group(1)} {month.group(2)} {month.group(3)}", "%B %d %Y"
        )
        return parsed.strftime("%m/%d/%Y")
    return _normalize_date(raw)


def _first(cells: dict[str, str], *needles: str) -> str:
    for key, value in cells.items():
        if any(needle in key for needle in needles) and value:
            return value
    return ""


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%m/%d/%Y")
    return str(value).strip()


def _digits(value: str) -> str | None:
    match = _PLAN_NUMBER.search(value or "")
    return match.group(1) if match else None


def _method(value: str) -> str | None:
    text = (value or "").upper()
    if "PRIOR" in text:
        return "PRIOR"
    if "CURRENT" in text:
        return "CURRENT"
    return None


def _yes_no(value: str) -> bool | None:
    text = (value or "").strip().lower()
    if text in {"yes", "y", "true"}:
        return True
    if text in {"no", "n", "false"}:
        return False
    return None


def _plan_type(value: str) -> str | None:
    text = (value or "").lower()
    if "403" in text:
        return "403b"
    if "401" in text:
        return "401k"
    return value or None


def _normalize_date(value: str) -> str | None:
    if not value:
        return None
    match = _DATE.search(value)
    if not match:
        return None
    try:
        parsed = datetime(int(match.group(3)), int(match.group(1)), int(match.group(2)))
        return parsed.strftime("%m/%d/%Y")
    except ValueError:
        return value
