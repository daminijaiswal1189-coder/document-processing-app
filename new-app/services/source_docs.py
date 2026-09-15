from __future__ import annotations

import fitz
from openpyxl import Workbook


SUMMARY_TEXT = """Compliance Testing Summary of Results
Plan Year End: December 31, 2025
MICHIGAN UNITED CREDIT UNION 401(K) SAVINGS PLAN
900062

416 Top Heavy Determination
Tested: 8.85% to Key Employees
Plan is not Top Heavy

410(b) Coverage Test
Standard: Pass Ratio percent 100%
401(k): Pass Ratio percent 100%

401(a)(4) Test
Tested: Pass

402(g) Deferral Limits
Tested: Pass

415(c) Annual Additions
Tested: Pass

Average Deferral Percentage (ADP)
Tested: Fail - Excess returns have already been processed
"""


def build_assembly_excel() -> bytes:
    """Val Assembly checklist in the layout of the attached Excel screenshot."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Val Assembly"
    headers = [
        "Omni Plan Number",
        "Plan Name",
        "PYE",
        "Plan Type",
        "Queue Picked Up By",
        "Rework",
        "Rush",
        "Per Payroll Match",
        "Is this a Re-Assemble",
        "If Re-Assemble Choose Reason",
        "If Assembly Error Choose Issue",
    ]
    values = [
        "800643",
        "DEERFIELD BEHAVIORAL HEALTH OF WARREN LLC 403(B) PLAN",
        "6/30/2026",
        "403(b)",
        "Shashikal",
        "No",
        "No",
        "Yes",
        "No",
        "",
        "",
    ]
    for col, header in enumerate(headers, start=1):
        sheet.cell(1, col, header)
        sheet.cell(2, col, values[col - 1])
    sheet["A5"] = "Additional Info:"
    extras = [
        (6, "If this is a 457 Plan is it a Non-Governmental plan? (Yes/No)", "No"),
        (7, "Do you have HCE Current Year? (yes/no)", "Yes"),
        (8, "Do you have HCE Future year? (yes/no)", "Yes"),
        (9, "Plan is Top heavy? (yes/no/403b n/a)", "403b N/A"),
        (10, "ADP/ACP Prior or current method test? (Prior / current / N/A SH)", "Current"),
        (
            11,
            "ADP/ACP Failure but All returns are reclassified (no returns needed)? (yes/no/N/A)",
            "N/A",
        ),
        (12, "Annual allocation (per document including true up)? (yes / no)", "No"),
        (13, "Notes (if needed)", ""),
    ]
    for row, label, value in extras:
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, value)
    buffer = __import__("io").BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def build_summary_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_textbox(fitz.Rect(54, 54, 558, 738), SUMMARY_TEXT, fontsize=11, fontname="helv")
    data = doc.tobytes()
    doc.close()
    return data
