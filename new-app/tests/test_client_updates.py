import io

import fitz
from openpyxl import Workbook

from services.file_order import location_index, order_uploads
from services.package_polish import (
    drop_brf_pages,
    exclude_brf_uploads,
    package_checks,
    remove_disallowed_wording,
    reorder_inner_pages,
)
from services.source_profile import parse_excel
from models.plan_profile import PlanProfile


def _pdf(pages: list[str]) -> bytes:
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def test_location_files_stay_with_their_section_in_number_order():
    uploads = [
        ("2025Census.pdf", b"%PDF"),
        ("Contribution Analysis Loc 2.pdf", b"%PDF"),
        ("Contribution Analysis Loc 1.pdf", b"%PDF"),
        ("2025ContriAnalysis.pdf", b"%PDF"),
    ]
    ordered = [name for name, _ in order_uploads(uploads)]
    assert ordered == [
        "2025ContriAnalysis.pdf",
        "Contribution Analysis Loc 1.pdf",
        "Contribution Analysis Loc 2.pdf",
        "2025Census.pdf",
    ]
    assert location_index("Contribution Analysis Loc 2.pdf") == 2


def test_brf_report_file_is_left_out_and_pages_are_dropped():
    kept, notes = exclude_brf_uploads(
        [("2025BRF.pdf", b"%PDF"), ("2025Census.pdf", b"%PDF")]
    )
    assert [name for name, _ in kept] == ["2025Census.pdf"]
    assert notes and "BRF" in notes[0]

    doc = fitz.open(stream=_pdf(["Benefits, Rights and Features Test\nPass", "Census\nPlan Number 111111"]), filetype="pdf")
    try:
        assert drop_brf_pages(doc)
        assert doc.page_count == 1
        assert "Census" in (doc[0].get_text("text") or "")
    finally:
        doc.close()


def test_inner_adp_and_410b_pages_follow_the_listed_order():
    doc = fitz.open(
        stream=_pdf(
            [
                "ADP/ACP Corrections Page\nRefunds",
                "ADP/ACP Test Detail\nRates",
                "ADP/ACP Summary Page\nResults",
                "410(b) Summary report\nPass",
                "410(b) Detail report\nEmployees",
            ]
        ),
        filetype="pdf",
    )
    try:
        notes = reorder_inner_pages(doc)
        text = [page.get_text("text") or "" for page in doc]
    finally:
        doc.close()
    assert notes
    assert text[0].startswith("ADP/ACP Summary")
    assert "Detail" in text[1]
    assert "Corrections" in text[2]
    assert "Detail" in text[3]
    assert "Summary" in text[4]


def test_disallowed_wording_is_removed():
    doc = fitz.open(
        stream=_pdf(
            [
                "Action Required\nEscalation notice for this failure\nForm 5500 is enclosed",
                "Statement of Contribution Report\nAmounts",
            ]
        ),
        filetype="pdf",
    )
    try:
        notes = remove_disallowed_wording(doc)
        text = "\n".join(page.get_text("text") or "" for page in doc)
    finally:
        doc.close()
    assert notes
    assert "Escalation" not in text
    assert "5500" not in text
    assert "Statement of Contribution" not in text
    assert "Action Required" in text


def test_assembly_log_follows_the_client_columns(tmp_path):
    from openpyxl import load_workbook

    from models.plan_profile import DetectedSection
    from services.assembly_log import append_row

    doc = fitz.open(
        stream=_pdf(
            [
                "Cover Letter\nPlan Number 800643",
                "Result Summary\nDEERFIELD BEHAVIORAL HEALTH 403(B) PLAN\n800643\n"
                "416 Top Heavy Determination\nTest: 8.85% to Key Employees\nPlan is not Top Heavy\n"
                "410(b) Coverage Test\nTest: Pass\n"
                "402(g) Deferral Limits\nTest: Fail\n"
                "Average Deferral Percentage (ADP)\nTest: Fail\n"
                "415(c) Annual Additions\nTest: Pass",
                "Future HCE\nSmith, Ann",
            ]
        ),
        filetype="pdf",
    )
    profile = PlanProfile(
        plan_number="800643",
        plan_name="Cover name should lose to the summary",
        hce_future=True,
        hce_current=False,
        contributions_required=True,
        assembly_notes="Annual allocation and true-up both has been done",
        fail_402g=True,
        fail_415=False,
        adp_failed=True,
        detected_sections=[DetectedSection(name="402(g)", page=2, matched_alias="402(g)")],
    )
    path = tmp_path / "ValAssemblyLog.xlsx"
    try:
        append_row(path, profile, "file.pdf", "job", doc=doc)
        append_row(
            path,
            PlanProfile(contributions_required=True, assembly_notes="true-up done, no variance"),
            "file.pdf",
            "job2",
        )
    finally:
        doc.close()
    book = load_workbook(path)
    sheet = book.active
    assert sheet["A1"].value == "Omni Plan Number"
    assert sheet["E2"].value == "Allocation Completed"
    assert str(sheet["A1"].fill.fgColor.rgb).endswith("808080")
    assert str(sheet["C1"].fill.fgColor.rgb).endswith("8EA9DB")
    assert str(sheet["U1"].fill.fgColor.rgb).endswith("808080")
    assert sheet["U1"].value == "Does this plan use New Comp"
    assert sheet.row_dimensions[1].height == 116
    assert sheet.row_dimensions[2].height == 58.5
    assert sheet.row_dimensions[3].height == 15.5
    assert sheet["A5"].value in (None, "")
    assert sheet.row_dimensions[5].height is None
    assert "A1:A2" in [str(item) for item in sheet.merged_cells.ranges]
    assert sheet["A3"].value == "800643"
    assert "DEERFIELD" in sheet["B3"].value
    assert sheet["C3"].value == "Yes"
    assert sheet["D3"].value == "No"
    assert sheet["E3"].value == "Annual-Var"
    assert sheet["F3"].value == "F"
    assert sheet["G3"].value == "F"
    assert sheet["H3"].value == "P"
    assert sheet["I3"].value == "N/A"
    assert sheet["L3"].value == "P"
    assert sheet["M3"].value == "No"
    assert sheet["O3"].value == "N/A"
    assert sheet["P3"].value in (None, "")
    assert sheet["E4"].value == "Variance"
    assert sheet["E4"].comment is not None
    assert "0.00 variance" in sheet["E4"].comment.text


def test_log_records_sent_date_and_build_time(tmp_path):
    from datetime import datetime

    from openpyxl import load_workbook

    from services.assembly_log import append_row

    path = tmp_path / "ValAssemblyLog.xlsx"
    append_row(
        path,
        PlanProfile(plan_number="100000", plan_name="Example Plan"),
        "file.pdf",
        "job",
        finished_at=datetime(2026, 10, 9, 17, 32),
        elapsed_seconds=12.4,
    )
    append_row(
        path,
        PlanProfile(plan_number="200000", plan_name="Second Plan"),
        "file.pdf",
        "job2",
        finished_at=datetime(2026, 10, 9, 17, 40),
        elapsed_seconds=75,
    )
    book = load_workbook(path)
    sheet = book.active
    assert sheet.cell(1, 21).value == "Does this plan use New Comp"
    assert sheet.cell(1, 22).value == "Date sent back to tester"
    assert sheet.cell(1, 23).value == "Time to build"
    assert sheet.cell(3, 22).value == "10/9/2026 5:32 PM"
    assert sheet.cell(3, 23).value == "12 sec"
    assert sheet.cell(4, 22).value == "10/9/2026 5:40 PM"
    assert sheet.cell(4, 23).value == "1 min 15 sec"
    assert sheet.cell(3, 21).value in (None, "")


def test_excel_include_contributions_and_prior_year():
    book = Workbook()
    sheet = book.active
    sheet["A5"] = "Additional Info:"
    sheet["A10"] = "ADP/ACP Prior or current method test? (Prior / current / N/A SH)"
    sheet["B10"] = "PRIOR YEAR"
    sheet["A12"] = "Annual allocation (per document including true up)? (yes / no)"
    sheet["B12"] = "Yes"
    sheet["A13"] = "Notes (if needed)"
    sheet["B13"] = "include contributions wordings"
    buffer = io.BytesIO()
    book.save(buffer)
    overlay = parse_excel(buffer.getvalue())
    assert overlay["testing_method"] == "PRIOR"
    assert overlay["contributions_required"] is True
    assert overlay["true_up"] is True


def test_result_summary_flags_top_heavy_mismatch():
    doc = fitz.open(
        stream=_pdf(
            [
                "Result Summary\n416 Top Heavy Determination\nTest: 70.00% to Key Employees\nPlan is not Top Heavy\n"
                "402(g) Deferral Limits\nTest: Pass\nPlan Number 111111\nSample Plan",
                "Top Heavy\nTop Heavy Percent: 55%\nPlan Number 111111\nSample Plan",
                "402(g) Test\nTest: Fail\nPlan Number 111111\nSample Plan",
            ]
        ),
        filetype="pdf",
    )
    try:
        items = {item.code: item for item in package_checks(doc, PlanProfile(plan_number="111111", plan_name="Sample Plan"))}
    finally:
        doc.close()
    assert items["top_heavy_summary"].passed is False
    assert "55" in items["top_heavy_summary"].detail
    assert items["summary_402(g)"].passed is False


def test_excess_notice_accepts_written_plan_year():
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Cover Letter\nMutual of America", fontsize=11)
    doc.new_page().insert_text(
        (72, 72),
        "Annual Valuation Report\n"
        "January 1, 2025 - December 31, 2025\n"
        "FAMILY HEALTH NETWORK OF CENTRAL NEW YORK INC. 401K PS PLAN\n"
        "801239\n"
        "Failed Compliance Testing - After 12-months\n"
        "Excess Return Notice\n"
        "Current method testing",
        fontsize=11,
    )
    try:
        items = {
            item.code: item
            for item in package_checks(
                doc,
                PlanProfile(
                    plan_number="801239",
                    plan_name="FAMILY HEALTH NETWORK OF CENTRAL NEW YORK INC. 401K PS PLAN",
                    plan_year_end="12/31/2025",
                ),
            )
        }
    finally:
        doc.close()
    assert items["report_identity"].passed is True


def test_variance_page_flags_zero_format_and_name_order():
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Variance Report\nPlan Number 111111\nSample Plan\n06/30/2026\n$0\nZed, Ann\nAnn, Bea", fontsize=11)
    page.insert_text((72, 180), "cut", fontsize=14)
    try:
        items = {item.code: item for item in package_checks(doc, PlanProfile(plan_number="111111", plan_name="Sample Plan", plan_year_end="06/30/2026"))}
    finally:
        doc.close()
    assert items["variance_client"].passed is False
    detail = items["variance_client"].detail
    assert "0.00" in detail
    assert "alphabetical" in detail
    assert "font size" in detail
